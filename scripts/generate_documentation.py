#!/usr/bin/env python3
"""Convert docs/ESTATEPILOT_AI_COMPLETE_GUIDE.md into a professional DOCX.

Usage:
  scripts/.venv/bin/python scripts/generate_documentation.py [--input FILE.md] [--output FILE.docx]

Supported Markdown subset: # / ## / ### / #### headings, paragraphs, bullet and numbered lists
(incl. "- [ ]" checkboxes), pipe tables, fenced code blocks (mermaid shown as code text), blockquotes,
**bold**, *italic*, `inline code`. Requires python-docx (see scripts/requirements.txt).
"""
import argparse
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is not installed.\nCreate an isolated environment and install it:\n"
             "  python3 -m venv scripts/.venv\n  scripts/.venv/bin/pip install -r scripts/requirements.txt\n"
             "  scripts/.venv/bin/python scripts/generate_documentation.py")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_IN = ROOT / "docs" / "ESTATEPILOT_AI_COMPLETE_GUIDE.md"
DEFAULT_OUT = ROOT / "docs" / "ESTATEPILOT_AI_COMPLETE_GUIDE.docx"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MONO = "Consolas"


# ---------- low-level helpers ----------
def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def add_field(paragraph, instr, placeholder=""):
    """Insert a complex field (PAGE, NUMPAGES, TOC ...)."""
    def fc(t):
        r = paragraph.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        if t == "begin":
            e.set(qn("w:dirty"), "true")
        r._r.append(e)
        return r
    fc("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    paragraph.add_run(placeholder)
    fc("end")


def set_cell_borders_and_margins(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "A6A6A6")
        borders.append(b)
    tblPr.append(borders)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_inline(paragraph, text, size=None, bold=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
            run.font.color.rgb = RGBColor(0x9C, 0x1C, 0x1C)
            if bold:
                run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            if bold:
                run.bold = True
        if size:
            run.font.size = Pt(size)
        if color is not None and not (part.startswith("`")):
            run.font.color.rgb = color


def split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip().replace("\\|", "|") for c in cells]


# ---------- document setup ----------
def setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before in (("Heading 1", 20, 18), ("Heading 2", 15, 14), ("Heading 3", 12.5, 10), ("Heading 4", 11, 8)):
        h = st[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = ACCENT
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.0)


def header_footer(section, title):
    section.different_first_page_header_footer = True  # no header/footer on cover
    hp = section.header.paragraphs[0]
    hp.text = title
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in hp.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    pPr = hp._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3A5F")):
        bt.set(qn(f"w:{k}"), v)
    bdr.append(bt)
    pPr.append(bdr)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    add_field(fp, "PAGE", "1")
    fp.add_run(" of ")
    add_field(fp, "NUMPAGES", "1")
    for r in fp.runs:
        r.font.size = Pt(9)


def cover(doc, title, subtitle_lines, status_note):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    r = p.add_run("AI-assisted lead management for real-estate agencies")
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    doc.add_paragraph()
    for line in subtitle_lines:
        p = doc.add_paragraph()
        add_inline(p, line, size=12)
    p = doc.add_paragraph()
    add_inline(p, f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC", size=11)
    if status_note:
        doc.add_paragraph()
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        c = t.rows[0].cells[0]
        shade(c._tc.get_or_add_tcPr(), "FFF4E5")
        c.paragraphs[0].text = ""
        add_inline(c.paragraphs[0], status_note, size=10)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc):
    h = doc.add_paragraph()
    r = h.add_run("Contents")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u',
              "Right-click here and choose 'Update Field' (or press F9) to build the table of contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    # ask Word to refresh fields on open
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


# ---------- block renderers ----------
def code_block(doc, lines, lang):
    if lang == "mermaid":
        lines = ["-- Mermaid diagram source (render at https://mermaid.live or in a Markdown viewer) --"] + lines
    for i, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.2)
        pf.keep_together = True
        shade(p._p.get_or_add_pPr(), "F2F2F2")
        run = p.add_run(line if line else " ")
        run.font.name = MONO
        run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
        run.font.size = Pt(8)
        if lang == "mermaid" and i == 0:
            run.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table_block(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    set_cell_borders_and_margins(t)
    for i, htxt in enumerate(header):
        c = t.rows[0].cells[i]
        c.paragraphs[0].text = ""
        add_inline(c.paragraphs[0], htxt, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c._tc.get_or_add_tcPr(), "1F3A5F")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        row = (row + [""] * ncols)[:ncols]
        for i, txt in enumerate(row):
            cells[i].paragraphs[0].text = ""
            add_inline(cells[i].paragraphs[0], txt, size=9)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
            if ri % 2 == 1:
                shade(cells[i]._tc.get_or_add_tcPr(), "F5F7FA")
    for row in t.rows:  # don't split rows across pages
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def list_item(doc, text, ordered, num=None):
    box = re.match(r"\[( |x|X)\]\s+(.*)", text)
    if box:
        text = ("☑ " if box.group(1).lower() == "x" else "☐ ") + box.group(2)
    if ordered:  # manual numbers: Word's "List Number" style would not restart numbering per list
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run(f"{num}.\t")
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    else:
        p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)


def quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    shade(p._p.get_or_add_pPr(), "EAF1FB")
    add_inline(p, text)
    for r in p.runs:
        r.italic = True


# ---------- main conversion ----------
def convert(md_text, out_path):
    lines = md_text.splitlines()
    doc = Document()
    setup_styles(doc)
    title = "EstatePilot AI"
    i = 0
    # cover data: first H1, following non-empty lines up to first '## ' (excluding the blockquote status)
    sub, status = [], ""
    if lines and lines[0].startswith("# "):
        title = lines[0][2:].strip()
        i = 1
        while i < len(lines) and not lines[i].startswith("## "):
            s = lines[i].strip()
            if s.startswith(">"):
                status = s.lstrip("> ").strip()
            elif s:
                sub.append(s)
            i += 1
    cover(doc, title, sub, status)
    header_footer(doc.sections[0], f"{title}  |  Complete Guide")
    toc(doc)

    stats = {"headings": 0, "tables": 0, "code_blocks": 0}
    n = len(lines)
    while i < n:
        line = lines[i]
        if line.startswith("```"):
            lang = line[3:].strip()
            i += 1
            buf = []
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            code_block(doc, buf, lang)
            stats["code_blocks"] += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            level = len(m.group(1))  # "##" chapter -> Heading 1, "###" -> Heading 2, "####" -> Heading 3
            text = re.sub(r"[*`]", "", m.group(2))
            h = doc.add_heading(text, level=max(1, min(level - 1, 3)))
            if level == 2:
                h.paragraph_format.page_break_before = True
            stats["headings"] += 1
            i += 1
            continue
        if line.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|?\s*$", lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < n and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            table_block(doc, rows)
            stats["tables"] += 1
            continue
        m = re.match(r"^\s*[-*]\s+(.*)", line)
        if m:
            list_item(doc, m.group(1), False)
            i += 1
            continue
        m = re.match(r"^\s*(\d+)\.\s+(.*)", line)
        if m:
            list_item(doc, m.group(2), True, m.group(1))
            i += 1
            continue
        if line.startswith(">"):
            quote(doc, line.lstrip("> ").strip())
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        # paragraph: join consecutive plain lines
        buf = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,4}\s|```|\||>|\s*[-*]\s|\s*\d+\.\s)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))
    doc.core_properties.title = f"{title} - Complete Guide"
    doc.core_properties.author = "EstatePilot AI documentation generator"
    doc.core_properties.subject = "Technical documentation"
    doc.save(out_path)
    return stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=str(DEFAULT_IN))
    ap.add_argument("--output", default=str(DEFAULT_OUT))
    a = ap.parse_args()
    src = Path(a.input)
    if not src.exists():
        sys.exit(f"Input Markdown not found: {src}")
    stats = convert(src.read_text(encoding="utf-8"), a.output)
    print(f"Wrote {a.output}: {stats['headings']} headings, {stats['tables']} tables, {stats['code_blocks']} code blocks")


if __name__ == "__main__":
    main()
